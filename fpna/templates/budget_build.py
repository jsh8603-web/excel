"""fpna.templates.budget_build — 예산·인건비 수립(부서별 인원·비용)."""
from __future__ import annotations

from dataclasses import dataclass, field

import fpna._bootstrap  # noqa: F401
import openpyxl
from openpyxl.utils import get_column_letter

from fpna import house_style as hs
from fpna.templates.base import QCReport, qc_no_formula_errors, qc_totals

TYPE = "budget_build"


@dataclass
class DeptLine:
    dept: str
    headcount: int
    avg_cost: float        # 1인당 연간 인건비


@dataclass
class BudgetInput:
    title: str = "예산·인건비 수립"
    subtitle: str = ""
    unit: str = "₩mn"
    depts: list = field(default_factory=list)   # list[DeptLine]


def golden_sample() -> BudgetInput:
    return BudgetInput(
        title="예산·인건비 수립 (골든샘플)", subtitle="구조 검증용 — 더미", unit="₩mn",
        depts=[DeptLine("영업", 10, 60), DeptLine("개발", 20, 80),
               DeptLine("관리", 5, 55)],
    )


def build(data: BudgetInput, *, mode="create", base_path=None) -> openpyxl.Workbook:
    wb = openpyxl.Workbook(); ws = wb.active
    ws.title = hs.safe_sheet_title("Budget")
    last_col = 4
    hs.set_widths(ws, {1: 20, 2: 12, 3: 16, 4: 16})
    r = hs.report_frame(ws, data.title, subtitle=data.subtitle,
                        unit=data.unit, last_col=last_col)

    for j, h in enumerate(["부서 (단위: %s)" % data.unit, "인원", "1인당 비용", "인건비 합"], 1):
        hs.set_cell(ws, r, j, h, role="header", align=hs.LEFT if j == 1 else hs.CENTER)
    r += 1
    data_start = r
    for d in data.depts:
        hs.set_cell(ws, r, 1, d.dept, role="label", align=hs.LEFT)
        hs.set_cell(ws, r, 2, d.headcount, role="input", number_format=hs.FMT_INT)
        hs.set_cell(ws, r, 3, d.avg_cost, role="input", number_format=hs.FMT_INT)
        hs.set_cell(ws, r, 4, "=B%d*C%d" % (r, r), role="calc", number_format=hs.FMT_INT)
        r += 1
    data_end = r - 1
    # 합계
    hs.set_cell(ws, r, 1, "합계", role="total", align=hs.LEFT, bold=True)
    hs.set_cell(ws, r, 2, "=SUM(B%d:B%d)" % (data_start, data_end), role="calc",
                number_format=hs.FMT_INT, bold=True)
    hs.set_cell(ws, r, 4, "=SUM(D%d:D%d)" % (data_start, data_end), role="calc",
                number_format=hs.FMT_INT, bold=True)
    for j in range(1, last_col + 1):
        ws.cell(row=r, column=j).border = hs.BORDER_TOP_STRONG
    hs.report_footer(ws, r + 2, source="인사 정원 · 인건비 단가표",
                     prepared_by="FP&A", last_col=last_col)
    return wb


def qc(wb, data: BudgetInput) -> QCReport:
    rep = QCReport(TYPE)
    qc_no_formula_errors(wb, rep)
    total = sum(d.headcount * d.avg_cost for d in data.depts)
    hc = sum(d.headcount for d in data.depts)
    rep.add("인원 합 > 0", hc > 0, "hc=%d" % hc)
    rep.add("인건비 합 계산", total >= 0, "합=%.0f" % total)
    rep.add("단위 표기", bool(data.unit))
    return rep


__all__ = ["TYPE", "DeptLine", "BudgetInput", "golden_sample", "build", "qc"]
